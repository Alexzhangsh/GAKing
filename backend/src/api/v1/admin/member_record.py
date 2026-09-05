# @ai-generated
"""
X02-1 用户会员记录查询与导出 API
路由前缀：/api/v1/admin/member-record
权限码：member:view（查询）/ member:export（导出）

接口清单：
1. GET  /api/v1/admin/member-record/records    会员记录列表（分页+用户/状态/套餐/时间筛选）
2. GET  /api/v1/admin/member-record/stats      会员数据统计（生效中/已到期/总数）
3. POST /api/v1/admin/member-record/export     会员记录导出（Excel）
"""
import logging
import os
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, Query, Request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from src.api.v1.response_util import (
    error_response,
    get_request_id,
    handle_service_exception,
    success_response,
)
from src.common.auth_util import require_any_permission
from src.config.x02_1_constants import (
    MEMBER_EXPORT_MAX_ROWS,
    MEMBER_STATUS_LABELS,
    PERM_MEMBER_EXPORT,
    PERM_MEMBER_VIEW,
)
from src.dao.member_package_dao import MemberPackageDAO
from src.dao.user_member_record_dao import UserMemberRecordDAO
from src.db.init_db import DatabaseManager
from src.models.business.user_member_record_model import UserMemberRecord

logger = logging.getLogger("api.admin.member_record")

router = APIRouter(prefix="/api/v1/admin/member-record", tags=["后台-会员记录(X02-1)"])

# 导出文件目录（相对 backend 运行目录）
EXPORT_FILE_DIR = "exports/member"


# ── 依赖注入 ──────────────


async def get_view_admin(
    payload: dict = Depends(require_any_permission([PERM_MEMBER_VIEW])),
) -> dict:
    return payload


async def get_export_admin(
    payload: dict = Depends(require_any_permission([PERM_MEMBER_EXPORT])),
) -> dict:
    return payload


# ── Schema ──────────────


class MemberRecordExportRequest(BaseModel):
    """会员记录导出请求"""
    user_id: Optional[int] = Field(None, description="平台用户ID筛选")
    status: Optional[str] = Field(None, max_length=16, description="会员状态筛选")
    package_id: Optional[int] = Field(None, description="套餐ID筛选")
    keyword: Optional[str] = Field(None, max_length=64, description="套餐名称关键字")
    start_time: Optional[str] = Field(None, description="开通时间起始（YYYY-MM-DD HH:MM:SS）")
    end_time: Optional[str] = Field(None, description="开通时间截止（YYYY-MM-DD HH:MM:SS）")


# ── 序列化 ──────────────


def _serialize_record(r: UserMemberRecord) -> Dict[str, Any]:
    return {
        "id": r.id,
        "user_id": r.user_id,
        "package_id": r.package_id,
        "package_name": r.package_name or "",
        "member_commission_rate": float(r.member_commission_rate) if r.member_commission_rate is not None else 0.0,
        "status": r.status,
        "status_label": MEMBER_STATUS_LABELS.get(r.status, r.status),
        "started_at": r.started_at.strftime("%Y-%m-%d %H:%M:%S") if r.started_at else "",
        "expire_at": r.expire_at.strftime("%Y-%m-%d %H:%M:%S") if r.expire_at else "",
        "order_id": r.order_id,
        "remark": r.remark or "",
        "create_time": r.create_time.strftime("%Y-%m-%d %H:%M:%S") if r.create_time else "",
    }


def _parse_time(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")


# ════════════════════════════════════════════════════════════
# 1. 会员记录列表
# ════════════════════════════════════════════════════════════


@router.get("/records")
async def list_records(
    request: Request,
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(20, ge=1, le=100, description="每页条数"),
    user_id: Optional[int] = Query(None, description="平台用户ID筛选"),
    status: Optional[str] = Query(None, max_length=16, description="会员状态筛选"),
    package_id: Optional[int] = Query(None, description="套餐ID筛选"),
    keyword: Optional[str] = Query(None, max_length=64, description="套餐名称关键字"),
    start_time: Optional[str] = Query(None, description="开通时间起始"),
    end_time: Optional[str] = Query(None, description="开通时间截止"),
    admin_user_id: int = Depends(get_view_admin),
):
    """会员记录列表（分页 + 多条件筛选）"""
    request_id = get_request_id(request)
    try:
        start_dt = _parse_time(start_time)
        end_dt = _parse_time(end_time)
        if start_dt and end_dt and start_dt >= end_dt:
            raise ValueError("起始时间必须早于截止时间")

        async with DatabaseManager.get_session() as session:
            dao = UserMemberRecordDAO(session)
            items, total = await dao.paginate_records(
                page=page, page_size=page_size,
                user_id=user_id, status=status, package_id=package_id,
                keyword=keyword, start_time=start_dt, end_time=end_dt,
            )
            data = {
                "total": total, "page": page, "page_size": page_size,
                "items": [_serialize_record(r) for r in items],
            }
            return success_response(data=data, request_id=request_id)
    except Exception as exc:
        return handle_service_exception(exc, request_id)


# ════════════════════════════════════════════════════════════
# 2. 会员数据统计
# ════════════════════════════════════════════════════════════


@router.get("/stats")
async def get_member_stats(
    request: Request,
    admin_user_id: int = Depends(get_view_admin),
):
    """会员数据统计（生效中/已到期/总记录数）"""
    request_id = get_request_id(request)
    try:
        async with DatabaseManager.get_session() as session:
            dao = UserMemberRecordDAO(session)
            active_items, active_total = await dao.paginate_records(page=1, page_size=1, status="active")
            expired_items, expired_total = await dao.paginate_records(page=1, page_size=1, status="expired")
            all_items, all_total = await dao.paginate_records(page=1, page_size=1)
            data = {
                "active_count": active_total,
                "expired_count": expired_total,
                "total_count": all_total,
            }
            return success_response(data=data, request_id=request_id)
    except Exception as exc:
        return handle_service_exception(exc, request_id)


# ════════════════════════════════════════════════════════════
# 3. 会员记录导出
# ════════════════════════════════════════════════════════════


def _create_excel_workbook(columns: List[tuple], rows: List[Dict[str, Any]]) -> Workbook:
    """创建 Excel 工作簿（与 B06-2 导出样式一致）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "会员记录"

    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="微软雅黑", size=10)
    cell_alignment = Alignment(vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, title in enumerate([c[1] for c in columns], 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, (field_name, _) in enumerate(columns, 1):
            value = row_data.get(field_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border

    for col_idx, (_, title) in enumerate(columns, 1):
        max_width = len(title) * 2
        for row_idx in range(2, len(rows) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                char_len = sum(2 if ord(c) > 127 else 1 for c in str(cell_value))
                max_width = max(max_width, char_len)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 4, 60)

    ws.freeze_panes = "A2"
    return wb


def _format_export_row(r: UserMemberRecord) -> Dict[str, Any]:
    return {
        "id": r.id,
        "user_id": r.user_id,
        "package_name": r.package_name or "",
        "member_commission_rate": f"{float(r.member_commission_rate or 0) * 100:.2f}%",
        "status": MEMBER_STATUS_LABELS.get(r.status, r.status or ""),
        "started_at": r.started_at.strftime("%Y-%m-%d %H:%M:%S") if r.started_at else "",
        "expire_at": r.expire_at.strftime("%Y-%m-%d %H:%M:%S") if r.expire_at else "",
        "order_id": r.order_id or "",
        "remark": r.remark or "",
        "create_time": r.create_time.strftime("%Y-%m-%d %H:%M:%S") if r.create_time else "",
    }


EXPORT_COLUMNS = [
    ("id", "记录ID"),
    ("user_id", "用户ID"),
    ("package_name", "套餐名称"),
    ("member_commission_rate", "会员分佣比例"),
    ("status", "会员状态"),
    ("started_at", "开通时间"),
    ("expire_at", "到期时间"),
    ("order_id", "开通订单ID"),
    ("remark", "备注"),
    ("create_time", "记录创建时间"),
]


@router.post("/export")
async def export_records(
    request: Request,
    body: MemberRecordExportRequest,
    admin_user_id: int = Depends(get_export_admin),
):
    """会员记录导出（Excel，单次最大 1000 行）"""
    request_id = get_request_id(request)
    try:
        start_dt = _parse_time(body.start_time)
        end_dt = _parse_time(body.end_time)
        if start_dt and end_dt and start_dt >= end_dt:
            raise ValueError("起始时间必须早于截止时间")

        async with DatabaseManager.get_session() as session:
            dao = UserMemberRecordDAO(session)
            # 先统计总数，超限拒绝
            _, total = await dao.paginate_records(
                page=1, page_size=1,
                user_id=body.user_id, status=body.status, package_id=body.package_id,
                keyword=body.keyword, start_time=start_dt, end_time=end_dt,
            )
            if total > MEMBER_EXPORT_MAX_ROWS:
                raise ValueError(
                    f"导出数据量({total}行)超过单次最大限制({MEMBER_EXPORT_MAX_ROWS}行)，请缩小筛选范围"
                )
            if total == 0:
                return error_response(msg="无符合条件的会员记录", request_id=request_id)

            records = await dao.list_for_export(
                user_id=body.user_id, status=body.status, package_id=body.package_id,
                keyword=body.keyword, start_time=start_dt, end_time=end_dt,
                limit=MEMBER_EXPORT_MAX_ROWS, offset=0,
            )

        rows = [_format_export_row(r) for r in records]
        wb = _create_excel_workbook(EXPORT_COLUMNS, rows)

        export_dir = os.path.join(os.getcwd(), EXPORT_FILE_DIR)
        os.makedirs(export_dir, exist_ok=True)
        file_name = f"member_record_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8]}.xlsx"
        file_path = os.path.join(export_dir, file_name)
        wb.save(file_path)
        file_size = os.path.getsize(file_path)

        return success_response(
            data={
                "file_name": file_name,
                "file_path": file_path,
                "file_size": file_size,
                "row_count": len(records),
            },
            msg=f"会员记录导出成功，共 {len(records)} 条",
            request_id=request_id,
        )
    except ValueError as ve:
        return error_response(msg=str(ve), request_id=request_id)
    except Exception as exc:
        return handle_service_exception(exc, request_id)
